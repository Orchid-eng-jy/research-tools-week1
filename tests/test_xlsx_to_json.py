import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "code"))

from xlsx_to_json import ConversionError, convert_workbook


class XlsxToJsonTests(unittest.TestCase):
	def setUp(self):
		self.temporary_directory = tempfile.TemporaryDirectory()
		self.directory = Path(self.temporary_directory.name)

	def tearDown(self):
		self.temporary_directory.cleanup()

	def _save_workbook(self, rows, name="input.xlsx"):
		path = self.directory / name
		workbook = Workbook()
		worksheet = workbook.active
		worksheet.title = "数据"
		for row in rows:
			worksheet.append(row)
		workbook.save(path)
		workbook.close()
		return path

	def test_converts_data_and_calculates_statistics(self):
		input_path = self._save_workbook(
			[
				["姓名", "分数", "日期", "通过"],
				["张三", 80, date(2026, 9, 1), True],
				["李四", 90, date(2026, 9, 2), False],
				["王五", None, None, True],
			]
		)
		output_path = self.directory / "output.json"

		payload = convert_workbook(input_path, output_path)

		self.assertEqual(payload["sheet"], "数据")
		self.assertEqual(payload["statistics"]["row_count"], 3)
		self.assertEqual(payload["statistics"]["column_count"], 4)
		self.assertEqual(payload["statistics"]["empty_cell_count"], 2)
		self.assertEqual(payload["statistics"]["columns"]["分数"]["numeric"]["average"], 85)
		self.assertEqual(payload["data"][0]["日期"], "2026-09-01T00:00:00")
		self.assertEqual(json.loads(output_path.read_text(encoding="utf-8")), payload)

	def test_skips_completely_blank_data_rows(self):
		input_path = self._save_workbook([["编号", "值"], [1, 10], [None, None], [2, 20]])

		payload = convert_workbook(input_path, self.directory / "output.json")

		self.assertEqual(payload["statistics"]["row_count"], 2)
		self.assertEqual([row["编号"] for row in payload["data"]], [1, 2])

	def test_rejects_duplicate_headers(self):
		input_path = self._save_workbook([["编号", "编号"], [1, 2]])

		with self.assertRaisesRegex(ConversionError, "表头不能重复"):
			convert_workbook(input_path, self.directory / "output.json")

	def test_can_select_a_named_sheet(self):
		input_path = self.directory / "multiple.xlsx"
		workbook = Workbook()
		workbook.active.append(["无关列"])
		selected = workbook.create_sheet("目标")
		selected.append(["项目", "数量"])
		selected.append(["A", 3])
		workbook.save(input_path)
		workbook.close()

		payload = convert_workbook(input_path, self.directory / "output.json", "目标")

		self.assertEqual(payload["sheet"], "目标")
		self.assertEqual(payload["data"], [{"项目": "A", "数量": 3}])


if __name__ == "__main__":
	unittest.main()
