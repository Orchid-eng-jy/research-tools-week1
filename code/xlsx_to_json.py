"""Convert a two-dimensional XLSX table to JSON and calculate statistics."""

from __future__ import annotations

import argparse
import json
import math
import sys
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


class ConversionError(ValueError):
	"""Raised when the workbook cannot be represented as a table."""


def _json_value(value: Any) -> Any:
	if isinstance(value, (datetime, date, time)):
		return value.isoformat()
	if isinstance(value, float) and not math.isfinite(value):
		raise ConversionError("表格包含不能写入 JSON 的 NaN 或无穷大数值")
	return value


def _is_empty(value: Any) -> bool:
	return value is None or value == ""


def _trim_table(rows: Iterable[tuple[Any, ...]]) -> list[list[Any]]:
	table = [list(row) for row in rows]
	while table and all(_is_empty(value) for value in table[-1]):
		table.pop()

	if not table:
		raise ConversionError("工作表为空")

	last_column = max(
		(index for row in table for index, value in enumerate(row) if not _is_empty(value)),
		default=-1,
	)
	if last_column < 0:
		raise ConversionError("工作表为空")

	return [row[: last_column + 1] + [None] * max(0, last_column + 1 - len(row)) for row in table]


def _headers(header_row: list[Any]) -> list[str]:
	headers: list[str] = []
	for column_number, value in enumerate(header_row, start=1):
		if _is_empty(value):
			raise ConversionError(f"第 {column_number} 列的表头为空")
		header = str(value).strip()
		if not header:
			raise ConversionError(f"第 {column_number} 列的表头为空")
		headers.append(header)

	duplicates = sorted(name for name, count in Counter(headers).items() if count > 1)
	if duplicates:
		raise ConversionError(f"表头不能重复：{', '.join(duplicates)}")
	return headers


def _value_type(value: Any) -> str:
	if _is_empty(value):
		return "empty"
	if isinstance(value, bool):
		return "boolean"
	if isinstance(value, (int, float)):
		return "number"
	if isinstance(value, (datetime, date, time)):
		return "date"
	return "string"


def _column_statistics(headers: list[str], rows: list[list[Any]]) -> dict[str, Any]:
	statistics: dict[str, Any] = {}
	for index, header in enumerate(headers):
		values = [row[index] for row in rows]
		type_counts = Counter(_value_type(value) for value in values)
		numbers = [value for value in values if isinstance(value, (int, float)) and not isinstance(value, bool)]
		column: dict[str, Any] = {
			"non_empty_count": len(values) - type_counts["empty"],
			"empty_count": type_counts["empty"],
			"type_counts": dict(sorted(type_counts.items())),
		}
		if numbers:
			column["numeric"] = {
				"min": min(numbers),
				"max": max(numbers),
				"sum": sum(numbers),
				"average": sum(numbers) / len(numbers),
			}
		statistics[header] = column
	return statistics


def convert_workbook(input_path: Path, output_path: Path, sheet_name: str | None = None) -> dict[str, Any]:
	"""Convert one worksheet and write the JSON document to output_path."""
	input_path = input_path.resolve()
	output_path = output_path.resolve()
	if not input_path.is_file():
		raise ConversionError(f"输入文件不存在：{input_path}")
	if input_path.suffix.lower() != ".xlsx":
		raise ConversionError("输入文件必须是 .xlsx 文件")
	if input_path == output_path:
		raise ConversionError("输入文件和输出文件不能相同")

	try:
		workbook = load_workbook(input_path, read_only=True, data_only=True)
	except Exception as exc:
		raise ConversionError(f"无法读取工作簿：{exc}") from exc

	try:
		if sheet_name is not None:
			if sheet_name not in workbook.sheetnames:
				raise ConversionError(
					f"工作表不存在：{sheet_name}；可用工作表：{', '.join(workbook.sheetnames)}"
				)
			worksheet = workbook[sheet_name]
		else:
			worksheet = workbook.active

		table = _trim_table(worksheet.iter_rows(values_only=True))
		headers = _headers(table[0])
		data_rows = [row for row in table[1:] if not all(_is_empty(value) for value in row)]
		json_rows = [
			{header: _json_value(value) for header, value in zip(headers, row)}
			for row in data_rows
		]
		empty_cells = sum(_is_empty(value) for row in data_rows for value in row)
		payload = {
			"source": input_path.name,
			"sheet": worksheet.title,
			"statistics": {
				"row_count": len(data_rows),
				"column_count": len(headers),
				"empty_cell_count": empty_cells,
				"non_empty_cell_count": len(data_rows) * len(headers) - empty_cells,
				"columns": _column_statistics(headers, data_rows),
			},
			"data": json_rows,
		}
	finally:
		workbook.close()

	output_path.parent.mkdir(parents=True, exist_ok=True)
	with output_path.open("w", encoding="utf-8") as output_file:
		json.dump(payload, output_file, ensure_ascii=False, indent=2, allow_nan=False)
		output_file.write("\n")
	return payload


def _parser() -> argparse.ArgumentParser:
	parser = argparse.ArgumentParser(description="将二维表 XLSX 文件转换为带统计信息的 JSON 文件")
	parser.add_argument("input", type=Path, help="输入的 .xlsx 文件")
	parser.add_argument("output", type=Path, nargs="?", help="输出 JSON 文件，默认与输入文件同名")
	parser.add_argument("--sheet", help="工作表名称，默认使用活动工作表")
	return parser


def main(argv: list[str] | None = None) -> int:
	args = _parser().parse_args(argv)
	output_path = args.output or args.input.with_suffix(".json")
	try:
		payload = convert_workbook(args.input, output_path, args.sheet)
	except ConversionError as exc:
		print(f"错误：{exc}", file=sys.stderr)
		return 1

	stats = payload["statistics"]
	print(
		f"转换完成：{output_path}（工作表 {payload['sheet']}，"
		f"{stats['row_count']} 行，{stats['column_count']} 列）"
	)
	return 0


if __name__ == "__main__":
	raise SystemExit(main())
